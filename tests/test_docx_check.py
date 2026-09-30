import datetime as dt
import os
import tempfile
import unittest

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from src.check import docx_check, spelling, spelling_ai
from src.check.docx_model import read_docx
from src.check.docx_structure import analyze

TODAY = dt.date(2026, 9, 30)


def _style_normal(document, size=14):
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    fmt = normal.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(6), Pt(6), 1.0
    fmt.first_line_indent = Cm(1)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _plain(paragraph, size=None, bold=None, italic=None, center=False, right=False):
    fmt = paragraph.paragraph_format
    fmt.first_line_indent = Cm(0)
    fmt.space_before = fmt.space_after = Pt(0)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        if italic is not None:
            run.italic = italic


def build(path, *, body=None, size=14, margins=(20, 20, 30, 20), number="Số:    /SKHCN-CĐS", title=None,
          date="Lào Cai, ngày      tháng 9 năm 2026", kinh_gui=True, party=False, letter=False, signer_size=13,
          soft_break_header=False):
    document = docx.Document()
    _style_normal(document, size)
    section = document.sections[0]
    section.page_width, section.page_height = (Mm(216), Mm(279)) if letter else (Mm(210), Mm(297))
    section.top_margin, section.bottom_margin, section.left_margin, section.right_margin = (Mm(m) for m in margins)
    header = document.add_table(rows=1, cols=2)
    left, right = header.rows[0].cells
    left.paragraphs[0].text = "ĐẢNG BỘ SỞ" if party else "UBND TỈNH LÀO CAI"
    _plain(left.paragraphs[0], 13, False, center=True)
    p = left.add_paragraph("SỞ KHOA HỌC VÀ CÔNG NGHỆ")
    _plain(p, 13, True, center=True)
    p = left.add_paragraph(number)
    _plain(p, 13, center=True)
    if not title:
        p = left.add_paragraph("V/v triển khai nhiệm vụ chuyển đổi số")
        _plain(p, 12, center=True)
    if party:
        right.paragraphs[0].text = "ĐẢNG CỘNG SẢN VIỆT NAM"
        _plain(right.paragraphs[0], 15, True, center=True)
    elif soft_break_header:
        cell = right.paragraphs[0]
        run = cell.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
        run.font.size, run.bold = Pt(12), True
        run.add_break()
        run2 = cell.add_run("Độc lập - Tự do - Hạnh phúc")
        run2.font.size, run2.bold = Pt(13), True
        _plain(cell, center=True)
    else:
        right.paragraphs[0].text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"
        _plain(right.paragraphs[0], 12, True, center=True)
        p = right.add_paragraph("Độc lập - Tự do - Hạnh phúc")
        _plain(p, 13, True, center=True)
    p = right.add_paragraph(date)
    _plain(p, 14, False, True, center=True)
    if title:
        for text, bold in ((title, True), ("Về việc triển khai nhiệm vụ chuyển đổi số", True)):
            p = document.add_paragraph(text)
            _plain(p, 14, bold, center=True)
    if kinh_gui and not title:
        p = document.add_paragraph("Kính gửi: Các phòng, đơn vị trực thuộc Sở.")
        _plain(p, center=True)
    for text in body or [
        "Căn cứ Kế hoạch số 215/KH-UBND ngày 30/12/2025 của UBND tỉnh về chuyển đổi số;",
        "Sở Khoa học và Công nghệ đề nghị các phòng, đơn vị triển khai thực hiện các nhiệm vụ được giao.",
        "Đề nghị các đơn vị nghiêm túc thực hiện./.",
    ]:
        document.add_paragraph(text)
    closing = document.add_table(rows=1, cols=2)
    a, b = closing.rows[0].cells
    a.paragraphs[0].text = "Nơi nhận:"
    _plain(a.paragraphs[0], 12, True, True)
    for line in ("- Như trên;", "- Lưu: VT, CĐS."):
        p = a.add_paragraph(line)
        _plain(p, 11)
    b.paragraphs[0].text = "KT. GIÁM ĐỐC"
    _plain(b.paragraphs[0], signer_size, True, center=True)
    p = b.add_paragraph("PHÓ GIÁM ĐỐC")
    _plain(p, signer_size, True, center=True)
    p = b.add_paragraph("Nguyễn Văn An")
    _plain(p, signer_size, True, center=True)
    document.save(path)


def failing(report, group=None):
    return {(g["key"], c["label"], c["current"]) for g in report["groups"] for c in g["checks"]
            if c["status"] == "fail" and not c.get("hidden") and (group is None or g["key"] == group)}


class DocxCheckTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.directory.name, "du_thao.docx")

    def tearDown(self):
        self.directory.cleanup()

    def run_check(self, profile="skhcn", **kwargs):
        build(self.path, **kwargs)
        report, document, structure, issues = docx_check.run_format(self.path, profile, today=TODAY)
        return report, document, structure, issues

    def test_correct_document_passes_skhcn(self):
        report, document, structure, issues = self.run_check()
        self.assertEqual("CV", structure.doc_type)
        self.assertTrue(document.paragraphs[structure.start].text.startswith("Kính gửi"))
        self.assertTrue(document.paragraphs[structure.last_body].text.endswith("./."))
        self.assertEqual(set(), failing(report), failing(report))
        self.assertEqual([], [i for i in issues if i["severity"] == "fail"])

    def test_layout_and_body_errors_are_found(self):
        report, *_ = self.run_check(size=13, margins=(15, 20, 25, 20), letter=True)
        fails = failing(report)
        labels = {label for _, label, _ in fails}
        self.assertIn("Khổ giấy", labels)
        self.assertIn(("layout", "Lề trên", "15 mm"), fails)
        self.assertIn(("layout", "Lề trái", "25 mm"), fails)
        self.assertIn(("body", "Cỡ chữ", "13 pt"), fails)

    def test_nd30_accepts_ranges(self):
        report, *_ = self.run_check(profile="nd30", size=13, margins=(25, 20, 35, 15))
        self.assertEqual(set(), failing(report, "layout") | failing(report, "body"))

    def test_signer_size_and_end_mark(self):
        report, *_ = self.run_check(signer_size=14, body=["Nội dung chính của văn bản."])
        fails = failing(report)
        self.assertIn(("closing", "Chức vụ người ký · Cỡ chữ", "14 pt"), fails)
        self.assertIn("Kết thúc nội dung bằng ./.", {label for _, label, _ in fails})
        report, *_ = self.run_check(signer_size=13.5)
        self.assertEqual(set(), failing(report, "closing"))

    def test_reference_number_and_date(self):
        report, *_ = self.run_check(number="Số: 5/SKHCN-CSĐ", date="Lào Cai, ngày 3 tháng 1 năm 2026")
        fails = {(label, current) for _, label, current in failing(report, "reference")}
        self.assertIn(("Số nhỏ hơn 10 có số 0 phía trước", "5"), fails)
        self.assertIn(("Mã đơn vị soạn thảo", "CSĐ"), fails)
        self.assertIn(("Ngày nhỏ hơn 10 có số 0 phía trước", "3"), fails)
        self.assertIn(("Tháng 1, 2 có số 0 phía trước", "1"), fails)
        unit = next(c for g in report["groups"] for c in g["checks"] if c["label"] == "Mã đơn vị soạn thảo")
        self.assertIn("CĐS", unit["hint"])

    def test_type_code_must_match_title(self):
        report, _, structure, _ = self.run_check(title="KẾ HOẠCH", number="Số: 12/TTR-SKHCN", kinh_gui=False,
                                                 body=["Thực hiện chỉ đạo của UBND tỉnh, Sở xây dựng kế hoạch như sau./."])
        self.assertEqual("KH", structure.doc_type)
        fails = {(label, current) for _, label, current in failing(report, "reference")}
        self.assertIn(("Viết đúng ký hiệu loại văn bản", "TTR"), fails)
        self.assertIn(("Ký hiệu khớp tên loại văn bản", "TTr (Tờ trình)"), fails)

    def test_decision_starts_at_legal_basis(self):
        _, document, structure, _ = self.run_check(title="QUYẾT ĐỊNH", number="Số: 12/QĐ-SKHCN", kinh_gui=False)
        self.assertEqual("QĐ", structure.doc_type)
        self.assertTrue(document.paragraphs[structure.start].text.startswith("Căn cứ"))

    def test_soft_line_breaks_in_header_are_split(self):
        report, document, *_ = self.run_check(soft_break_header=True)
        roles = {p.role for p in document.paragraphs}
        self.assertTrue({"quoc_hieu", "tieu_ngu"} <= roles)
        self.assertEqual(set(), failing(report, "header"))

    def test_party_document_skips_nd30_header_rules(self):
        report, _, structure, _ = self.run_check(party=True, number="Số 12-CV/ĐU")
        self.assertTrue(structure.party)
        self.assertNotIn("header", {g["key"] for g in report["groups"]})

    def test_code_spelling(self):
        _, _, _, issues = self.run_check(body=[
            "Nâng cao năng xuất lao động và và triển khai ngiệp vụ.",
            "Ứng dụng internet, logic và KPI trong chuyển đổi số./."])
        found = {(i["category"], i["original"], i["suggestion"]) for i in issues}
        self.assertIn(("spelling", "năng xuất", "năng suất"), found)
        self.assertIn(("spelling", "ngiệp", "nghiệp"), found)
        self.assertIn(("repeat", "và và", "và"), found)
        self.assertFalse(any(i["original"] in ("internet", "logic", "KPI") for i in issues))

    def test_reduplication_and_compounds_are_not_repeats(self):
        issues = spelling.check_paragraph(_fake("Hai hệ thống chạy song song; xử lý hành vi vi phạm; "
                                                "liên kết kết quả; Luật Chuyển đổi số số 148/2025/QH15."), set())
        self.assertEqual([], [i for i in issues if i["category"] == "repeat"])

    def test_tone_placement_styles_are_both_valid(self):
        issues = spelling.check_paragraph(_fake("Hòa bình, hoà bình, thủy lợi, thuỷ lợi, kỹ thuật, kĩ thuật."), set())
        self.assertEqual([], [i for i in issues if i["category"] == "spelling"])


class AiFilterTests(unittest.TestCase):
    def test_ai_results_are_filtered(self):
        paragraph = _fake("Đề nghị triễn khai thực hiện theo Công văn số 368/KH-UBND./.", index=7)
        raw = [
            {"paragraph_id": "p7", "original": "triễn khai", "suggestion": "triển khai", "confidence": 0.95},
            {"paragraph_id": "p7", "original": "thực hiện", "suggestion": "thực hiện", "confidence": 0.9},
            {"paragraph_id": "p7", "original": "./.", "suggestion": ".", "confidence": 0.9},
            {"paragraph_id": "p7", "original": "368/KH-UBND", "suggestion": "368/KH-UBND.", "confidence": 0.9},
            {"paragraph_id": "p9", "original": "abc", "suggestion": "abd", "confidence": 0.9},
            {"paragraph_id": "p7", "original": "không có", "suggestion": "có", "confidence": 0.9},
            {"paragraph_id": "p7", "original": "Đề nghị", "suggestion": "Kính đề nghị", "confidence": 0.6,
             "explanation": "Không cần sửa"},
        ]
        kept, dropped = spelling_ai.filter_errors(raw, {7: paragraph}, [], last_index=7)
        self.assertEqual([("triễn khai", "triển khai", "fail")], [(k["original"], k["suggestion"], k["severity"]) for k in kept])
        self.assertEqual(6, dropped)


def _fake(text, index=1):
    from src.check.docx_model import Paragraph
    return Paragraph(index=index, text=text, style_id=None, style_name=None, in_table=False, table_index=None,
                     cell=None, section=0, jc="both", ind_left=0, ind_right=0, first_line=567, before=120, after=120,
                     line=240, line_rule="auto", contextual=False, numbered=False, role="body")


class ReaderTests(unittest.TestCase):
    def test_inherited_font_and_empty_spacing_tag(self):
        directory = tempfile.TemporaryDirectory()
        path = os.path.join(directory.name, "x.docx")
        document = docx.Document()
        _style_normal(document)
        paragraph = document.add_paragraph("Nội dung kế thừa định dạng từ style Normal.")
        paragraph._p.get_or_add_pPr().append(paragraph._p.get_or_add_pPr().makeelement(qn("w:spacing"), {}))
        section = document.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        document.save(path)
        parsed = read_docx(path)
        p = next(p for p in parsed.paragraphs if p.text.startswith("Nội dung"))
        self.assertEqual("Times New Roman", p.runs[0].font)
        self.assertEqual(14, p.main_size())
        self.assertEqual((120, 120, 240), (p.before, p.after, p.line))
        self.assertEqual("both", p.jc)
        self.assertEqual(567, p.first_line)
        self.assertIsNotNone(analyze(parsed))
        directory.cleanup()


if __name__ == "__main__":
    unittest.main()
